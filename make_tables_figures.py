#!/usr/bin/env python3
"""Tables_and_Figures.docx — Tables 1-6 (rendered) + Figures 1-7 (images + captions)."""
import json
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from PIL import Image

CAPS = json.load(open('/tmp/claude-0/-home-user-test/427d270c-c300-5c93-96fb-75f0dc166813/scratchpad/figbaa/captions.json'))
FIGDIR = '/tmp/claude-0/-home-user-test/427d270c-c300-5c93-96fb-75f0dc166813/scratchpad/figbaa'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, m, Cm(2.0))
USABLE_W = 17.0 / 2.54
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(10.5)


def title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    m = text.split(' ', 1)
    r = p.add_run(m[0]); r.bold = True
    if len(m) > 1:
        p.add_run(' ' + m[1])
    return p


def footnote(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text); r.font.size = Pt(9); r.italic = True
    return p


def make_table(headers, rows, widths=None, indent_flags=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hc = t.rows[0].cells
    for i, h in enumerate(headers):
        hc[i].text = ''
        pr = hc[i].paragraphs[0]; run = pr.add_run(h); run.bold = True; run.font.size = Pt(10)
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            pr = cells[i].paragraphs[0]
            run = pr.add_run(str(val)); run.font.size = Pt(10)
            if i == 0 and indent_flags and row is indent_flags.get('ref'):
                pass
            pr.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
    if widths:
        for i, w in enumerate(widths):
            for cell in t.columns[i].cells:
                cell.width = Cm(w)
    return t


def page():
    doc.add_page_break()


# =========================== TABLE 1 ===========================
title("Table 1. Baseline characteristics of the training and internal validation cohorts.")
h1 = ["Variable", "All (N=188)", "Training (N=131)", "Validation (N=57)", "P"]
# (label, all, train, valid, p)  — sub-rows use a leading indent for no/yes
r1 = [
 ["Massive haemoptysis", "", "", "", "0.929"],
 ["    no", "146 (77.66)", "101 (77.10)", "45 (78.95)", ""],
 ["    yes", "42 (22.34)", "30 (22.90)", "12 (21.05)", ""],
 ["Age (years)", "61.0 [54.0, 70.0]", "60.0 [53.0, 69.5]", "65.0 [55.0, 70.0]", "0.222"],
 ["Sex", "", "", "", "0.709"],
 ["    female", "58 (30.85)", "42 (32.06)", "16 (28.07)", ""],
 ["    male", "130 (69.15)", "89 (67.94)", "41 (71.93)", ""],
 ["Smoking status", "", "", "", "0.688"],
 ["    no", "75 (39.89)", "54 (41.22)", "21 (36.84)", ""],
 ["    yes", "113 (60.11)", "77 (58.78)", "36 (63.16)", ""],
 ["Duration of smoking (years)", "15.5 [0.0, 40.0]", "20.0 [0.0, 40.0]", "15.0 [0.0, 40.0]", "0.786"],
 ["Cancer", "", "", "", "0.185"],
 ["    no", "157 (83.51)", "113 (86.26)", "44 (77.19)", ""],
 ["    yes", "31 (16.49)", "18 (13.74)", "13 (22.81)", ""],
 ["Tuberculosis", "", "", "", "0.103"],
 ["    no", "159 (84.57)", "115 (87.79)", "44 (77.19)", ""],
 ["    yes", "29 (15.43)", "16 (12.21)", "13 (22.81)", ""],
 ["Bronchiectasis", "", "", "", "0.222"],
 ["    no", "118 (62.77)", "78 (59.54)", "40 (70.18)", ""],
 ["    yes", "70 (37.23)", "53 (40.46)", "17 (29.82)", ""],
 ["Fungal infection", "", "", "", "1.000"],
 ["    no", "165 (87.77)", "115 (87.79)", "50 (87.72)", ""],
 ["    yes", "23 (12.23)", "16 (12.21)", "7 (12.28)", ""],
 ["BAA", "", "", "", "0.049"],
 ["    no", "150 (79.79)", "110 (83.97)", "40 (70.18)", ""],
 ["    yes", "38 (20.21)", "21 (16.03)", "17 (29.82)", ""],
 ["Antiplatelet drugs", "", "", "", "0.757"],
 ["    no", "176 (93.62)", "123 (93.89)", "53 (92.98)", ""],
 ["    yes", "12 (6.38)", "8 (6.11)", "4 (7.02)", ""],
 ["Antithrombotic drugs", "", "", "", "1.000"],
 ["    no", "182 (96.81)", "127 (96.95)", "55 (96.49)", ""],
 ["    yes", "6 (3.19)", "4 (3.05)", "2 (3.51)", ""],
 ["PLT (10⁹/L)", "204.0 [177.5, 256.3]", "206.0 [184.0, 249.5]", "203.0 [169.0, 258.0]", "0.529"],
 ["D-dimer (mg/L)", "0.34 [0.20, 0.69]", "0.38 [0.19, 0.66]", "0.33 [0.23, 0.70]", "0.772"],
 ["INR", "0.98 [0.92, 1.05]", "0.98 [0.94, 1.03]", "0.98 [0.92, 1.07]", "0.828"],
 ["Fibrinogen (g/L)", "3.28 [2.51, 4.18]", "3.31 [2.46, 4.05]", "3.24 [2.60, 4.27]", "0.582"],
]
make_table(h1, r1, widths=[5.2, 3.3, 3.3, 3.3, 1.6])
footnote("Data are n (%) for categorical variables and median [Q1, Q3] for continuous variables. "
         "BAA, bronchial artery abnormality, defined as the presence of bronchial artery–pulmonary "
         "fistula (BAPF), bronchial artery malformation (BAM), or both (n=27 BAPF and n=18 BAM, with "
         "7 patients carrying both). PLT, platelet count; INR, international normalised ratio.")
page()

# =========================== TABLE 2 ===========================
title("Table 2. Univariable and multivariable logistic regression for massive haemoptysis (training cohort).")
h2 = ["Variable", "Univariable OR (95% CI)", "P", "Multivariable OR (95% CI)", "P"]
DASH = "—"
r2 = [
 ["Age", "1.02 (0.983–1.047)", "0.360", DASH, ""],
 ["Sex", "1.40 (0.563–3.461)", "0.472", DASH, ""],
 ["Smoking status", "0.89 (0.392–2.037)", "0.789", DASH, ""],
 ["Duration of smoking", "1.00 (0.977–1.018)", "0.813", DASH, ""],
 ["Cancer", "0.64 (0.171–2.368)", "0.501", DASH, ""],
 ["Tuberculosis", "4.23 (1.429–12.505)", "0.009", "5.18 (1.419–18.931)", "0.013"],
 ["Bronchiectasis", "2.36 (1.031–5.409)", "0.042", DASH, ""],
 ["Fungal infection", "1.14 (0.339–3.838)", "0.831", DASH, ""],
 ["BAA", "5.27 (1.960–14.162)", "0.001", "4.32 (1.325–14.098)", "0.015"],
 ["Antiplatelet drugs", "0.46 (0.055–3.921)", "0.480", DASH, ""],
 ["Antithrombotic drugs", "1.13 (0.113–11.244)", "0.919", DASH, ""],
 ["PLT", "1.00 (0.990–1.003)", "0.290", DASH, ""],
 ["D-dimer", "1.13 (0.920–1.380)", "0.247", DASH, ""],
 ["INR", "0.09 (0.001–7.670)", "0.284", DASH, ""],
 ["Fibrinogen", "0.57 (0.386–0.830)", "0.004", "0.65 (0.446–0.947)", "0.025"],
 ["Rad-score", "7.96 (2.745–23.105)", "<0.001", "4.73 (1.530–14.632)", "0.007"],
]
make_table(h2, r2, widths=[3.6, 4.4, 1.6, 4.4, 1.6])
footnote("OR, odds ratio; CI, confidence interval. Only variables retained in the multivariable "
         "model carry an adjusted estimate; — denotes not retained. Abbreviations as in Table 1; "
         "Rad-score, radiomics score.")
page()

# =========================== TABLE 3 ===========================
title("Table 3. Diagnostic performance of the clinical–radiomic nomogram.")
h3 = ["Cohort", "Accuracy", "Sensitivity", "Specificity", "AUC (95% CI)"]
r3 = [
 ["Training (n=131)", "0.824 (108/131)", "0.767 (23/30)", "0.842 (85/101)", "0.850 (0.775–0.925)"],
 ["Validation (n=57)", "0.684 (39/57)", "0.750 (9/12)", "0.667 (30/45)", "0.794 (0.631–0.958)"],
]
make_table(h3, r3, widths=[3.6, 3.3, 3.3, 3.3, 3.5])
footnote("Accuracy, sensitivity and specificity were obtained at an estimated-probability cut-point "
         "of 0.255, derived from the maximum Youden index in the training cohort and applied "
         "unchanged to the validation cohort. AUC, area under the receiver operating characteristic curve.")
page()

# =========================== TABLE 4 ===========================
title("Table 4. Distribution of the Rad-score by presentation group.")
h4 = ["Cohort", "Group", "n", "Median (Q1, Q3)", "P"]
r4 = [
 ["Training", "Non-massive haemoptysis", "101", "−1.361 (−1.640, −1.062)", "1.49×10⁻⁵"],
 ["", "Massive haemoptysis", "30", "−1.043 (−1.189, −0.706)", ""],
 ["Validation", "Non-massive haemoptysis", "45", "−1.189 (−1.523, −0.954)", "0.122"],
 ["", "Massive haemoptysis", "12", "−1.029 (−1.275, −0.596)", ""],
]
make_table(h4, r4, widths=[2.6, 5.4, 1.4, 5.6, 2.0])
footnote("Groups were compared with the Mann–Whitney U test. Q1, first quartile; Q3, third quartile.")
page()

# =========================== TABLE 5 ===========================
title("Table 5. Discrimination of the three models and pairwise comparison (DeLong test), by cohort.")
h5 = ["Cohort", "Comparison", "AUC (Model 1)", "AUC (Model 2)", "Z", "P"]
r5 = [
 ["Training", "Combined vs. Rad-score", "0.850", "0.761", "2.058", "0.040"],
 ["", "Combined vs. Clinical", "0.850", "0.796", "1.841", "0.066"],
 ["", "Rad-score vs. Clinical", "0.761", "0.796", "−0.514", "0.608"],
 ["Validation", "Combined vs. Rad-score", "0.794", "0.647", "1.373", "0.170"],
 ["", "Combined vs. Clinical", "0.794", "0.744", "0.967", "0.333"],
 ["", "Rad-score vs. Clinical", "0.647", "0.744", "−0.665", "0.506"],
]
make_table(h5, r5, widths=[2.4, 4.8, 2.9, 2.9, 1.8, 2.2])
footnote("Model AUCs with 95% CIs — training: combined 0.850 (0.775–0.925), clinical 0.796 "
         "(0.693–0.898), Rad-score 0.761 (0.672–0.850); validation: combined 0.794 "
         "(0.631–0.958), clinical 0.744 (0.555–0.934), Rad-score 0.647 (0.456–0.839). "
         "Combined model = Rad-score + tuberculosis + BAA + fibrinogen; clinical model = "
         "tuberculosis + BAA + fibrinogen. With three comparisons per cohort, a Bonferroni-corrected "
         "threshold is P=0.0167.")
page()

# =========================== TABLE 6 ===========================
title("Table 6. Continuous NRI and IDI of the combined model versus the clinical model.")
h6 = ["Cohort", "NRI (95% CI)", "P", "IDI (95% CI)", "P"]
r6 = [
 ["Training", "0.428 (0.029, 0.827)", "0.036", "0.062 (0.013, 0.111)", "0.013"],
 ["Validation", "0.544 (−0.076, 1.165)", "0.085", "0.064 (−0.044, 0.171)", "0.245"],
]
make_table(h6, r6, widths=[2.8, 4.4, 1.6, 4.4, 1.6])
footnote("NRI, continuous net reclassification improvement; IDI, integrated discrimination "
         "improvement, for adding the Rad-score to the clinical model. Continuous NRI rejects the "
         "null more often than its nominal level implies and should be read alongside the IDI.")
page()

# =========================== FIGURES 1-7 ===========================
def add_caption(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.line_spacing = 1.25
    m = text.split(' ', 1)
    r = p.add_run(m[0]); r.bold = True; r.font.size = Pt(10)
    if len(m) > 1:
        r2 = p.add_run(' ' + m[1]); r2.font.size = Pt(10)

for i in range(1, 8):
    img = f'{FIGDIR}/fig{i}.png'
    w, h = Image.open(img).size
    disp_w = USABLE_W
    disp_h = disp_w * h / w
    if disp_h > 20.0 / 2.54:
        disp_h = 20.0 / 2.54; disp_w = disp_h * w / h
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(img, width=Inches(disp_w))
    add_caption(CAPS[f'image{i}.tiff'])
    if i < 7:
        page()

out = '/home/user/test/Tables_and_Figures.docx'
doc.save(out)
print('saved', out)
